#!/usr/bin/env python3
"""Render the gap analysis as a Word document.

    python3 build_analysis.py analysis.md <client>_gap_analysis.docx

Author the analysis as markdown first — it is easier to draft and review — then render
it. Supported: # ## ### headings, paragraphs, - and * bullets, 1. numbered lists,
| pipe | tables | with a --- separator row, ``` fenced blocks, and **bold** / `code`
inline.

Writes .docx when python-docx is importable. Otherwise it copies the markdown beside
the requested path and says so, rather than failing silently.
"""
import os
import re
import shutil
import sys

ACCENT = "2E3B4E"


def _runs(paragraph, text):
    """Emit **bold** and `code` as runs; everything else plain."""
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            r.font.name = "Consolas"
        else:
            paragraph.add_run(part)


def _table(doc, block):
    from docx.shared import Pt
    rows = [r for r in block if not re.match(r'^\|[\s:|-]+\|$', r.strip())]
    cells = [[c.strip() for c in r.strip().strip('|').split('|')] for r in rows]
    if not cells:
        return
    width = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=width)
    t.style = "Table Grid"
    for ri, row in enumerate(cells):
        for ci in range(width):
            cell = t.cell(ri, ci)
            cell.text = ""
            p = cell.paragraphs[0]
            _runs(p, row[ci] if ci < len(row) else "")
            for run in p.runs:
                run.font.size = Pt(9)
                if ri == 0:
                    run.bold = True
    doc.add_paragraph()


def render(md_path, out_path):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for lvl, size in ((1, 20), (2, 14), (3, 11.5)):
        st = doc.styles["Heading %d" % lvl]
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(ACCENT)

    lines = open(md_path, encoding="utf-8").read().split("\n")
    i, pending = 0, []

    def flush_table():
        if pending:
            _table(doc, list(pending))
            pending.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("|"):
            pending.append(line)
            i += 1
            continue
        flush_table()

        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            r = p.add_run("\n".join(buf))
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            continue

        if not stripped:
            i += 1
            continue

        m = re.match(r'^(#{1,3})\s+(.*)$', stripped)
        if m:
            doc.add_heading(m.group(2).strip(), level=len(m.group(1)))
            i += 1
            continue

        m = re.match(r'^[-*]\s+(.*)$', stripped)
        if m:
            _runs(doc.add_paragraph(style="List Bullet"), m.group(1))
            i += 1
            continue

        m = re.match(r'^\d+\.\s+(.*)$', stripped)
        if m:
            _runs(doc.add_paragraph(style="List Number"), m.group(1))
            i += 1
            continue

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _runs(p, stripped)
        i += 1

    flush_table()
    doc.save(out_path)
    return out_path


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        return 2
    md_path, out_path = sys.argv[1], sys.argv[2]
    if not os.path.exists(md_path):
        print("no such file: %s" % md_path)
        return 2
    try:
        import docx  # noqa: F401
    except ImportError:
        fallback = os.path.splitext(out_path)[0] + ".md"
        if os.path.abspath(md_path) != os.path.abspath(fallback):
            shutil.copy(md_path, fallback)
        print("python-docx not installed -> left the analysis at %s" % fallback)
        print("For the Word document: use the session's document skill, or "
              "pip install python-docx.")
        return 0
    print("Wrote %s" % render(md_path, out_path))
    return 0


if __name__ == "__main__":
    sys.exit(main())
